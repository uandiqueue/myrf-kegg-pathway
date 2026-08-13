"""Round-2 literature and artifact audit for the curated MYRF pathway."""

from __future__ import annotations

from collections import Counter
from datetime import datetime
from hashlib import sha256
from pathlib import Path
import re
import sys
import xml.etree.ElementTree as ET

import openpyxl
import pandas as pd


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
OUTPUT = ROOT / "output"
WORKBOOK = ROOT / "raw_data" / "myrf-final-consolidation-documentation.xlsx"
AUDIT_CSV = DATA / "literature_audit.csv"
REPORT = OUTPUT / "literature_audit_round2.md"

MATCH_COLUMNS = [
    "source",
    "target",
    "primary_relationship",
    "kegg_edge_subtype",
    "direction",
    "regulation_level",
    "confidence",
]

EXPECTED_SUBTYPES = {
    "expression": 14,
    "indirect effect": 8,
    "repression": 7,
    "inhibition": 1,
    "binding/association": 1,
    "ubiquitination": 1,
}

FORBIDDEN_EDGES = {
    ("MAG", "MYRF_gene", "others/unknown"),
    ("FYN", "MYRF_gene", "expression"),
    ("MYRF_protein", "SOX10", "activation"),
}


def clean(value: object) -> str:
    if value is None or pd.isna(value):
        return ""
    return str(value).strip()


def bool_value(value: object) -> bool:
    return clean(value).lower() in {"true", "1", "yes"}


def file_hash(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def edge_tuple(row: pd.Series) -> tuple[str, ...]:
    return tuple(clean(row[col]) for col in MATCH_COLUMNS)


def write_report(
    checks: list[tuple[str, bool, str]],
    edges: pd.DataFrame,
    audit: pd.DataFrame,
    errors: list[str],
) -> None:
    passed = not errors and all(ok for _, ok, _ in checks)
    lines = [
        "# MYRF pathway literature audit — round 2",
        "",
        f"Generated: {datetime.now().astimezone().isoformat(timespec='seconds')}",
        "",
        f"Result: **{'PASS' if passed else 'FAIL'}**",
        "",
        "## Scope",
        "",
        "This second pass checks the curated edge register against the primary-literature "
        "audit table and then verifies that the workbook, CSVs, DOT/SVG, KGML, evidence "
        "export, and synchronized Pathview copy all encode the same pathway.",
        "",
        "## Automated checks",
        "",
        "| Check | Result | Detail |",
        "|---|---:|---|",
    ]
    for name, ok, detail in checks:
        lines.append(f"| {name} | {'PASS' if ok else 'FAIL'} | {detail} |")

    lines += [
        "",
        "## Final model",
        "",
        f"- Nodes: {len(pd.read_csv(DATA / 'final_nodes.csv'))}",
        f"- Edges: {len(edges)}",
        f"- Strong edges: {(edges['confidence'].str.lower() == 'strong').sum()}",
        f"- Moderate edges: {(edges['confidence'].str.lower() == 'moderate').sum()}",
        f"- Manual-review flags: {edges['needs_manual_review'].map(bool_value).sum()}",
        "",
        "Subtype counts: " + ", ".join(
            f"{name}={count}" for name, count in edges["kegg_edge_subtype"].value_counts().items()
        ),
        "",
        "## Literature-reviewed edge register",
        "",
        "| Source | Target | Semantics | Confidence | Primary source | Evidence scope |",
        "|---|---|---|---:|---|---|",
    ]
    audit_by_key = {edge_tuple(row): row for _, row in audit.iterrows()}
    for _, row in edges.iterrows():
        audited = audit_by_key[edge_tuple(row)]
        source_url = clean(audited["primary_source"])
        lines.append(
            f"| {row['source']} | {row['target']} | {row['kegg_edge_subtype']} "
            f"({row['direction']}; {row['regulation_level']}) | {row['confidence']} | "
            f"[primary source]({source_url}) | {clean(audited['evidence_summary'])} |"
        )

    lines += [
        "",
        "## Explicit uncertainty boundaries",
        "",
        "- MAG, ANGPTL2, and FYN are retained as moderate indirect effects on MYRF protein abundance; the terminal mechanism is unresolved.",
        "- MAPK1 and MAPK3 are retained separately for identifier compatibility, but the experiment supports combined ERK1/2 activity rather than resolved individual contributions.",
        "- SOX8 is an indirect moderate edge based on partial genetic redundancy with SOX10; direct MYRF enhancer binding was not shown.",
        "- MPZ and EGR2 repression is limited to heterologous reporter assays and is therefore moderate.",
        "",
        "## Removed or collapsed representations",
        "",
        "- Removed the duplicate MAG→MYRF unknown-mechanism edge.",
        "- Removed the unsupported direct FYN→MYRF transcriptional edge.",
        "- Collapsed directional SOX10↔MYRF protein activation into one undirected binding/association edge.",
        "- Excluded the duplicate PLP1 WmN1 row because that enhancer did not show MYRF–SOX10 synergy.",
    ]
    if errors:
        lines += ["", "## Failures", ""] + [f"- {error}" for error in errors]

    REPORT.parent.mkdir(parents=True, exist_ok=True)
    REPORT.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> int:
    errors: list[str] = []
    checks: list[tuple[str, bool, str]] = []

    def check(name: str, condition: bool, detail: str) -> None:
        checks.append((name, condition, detail))
        if not condition:
            errors.append(f"{name}: {detail}")

    edges = pd.read_csv(DATA / "final_edges.csv", dtype=str).fillna("")
    nodes = pd.read_csv(DATA / "final_nodes.csv", dtype=str).fillna("")
    audit = pd.read_csv(AUDIT_CSV, dtype=str).fillna("")
    evidence = pd.read_csv(OUTPUT / "myrf_pathway_evidence.csv", dtype=str).fillna("")

    check("edge count", len(edges) == 32, f"found {len(edges)}; expected 32")
    check("node count", len(nodes) == 31, f"found {len(nodes)}; expected 31")
    check(
        "subtype distribution",
        edges["kegg_edge_subtype"].value_counts().to_dict() == EXPECTED_SUBTYPES,
        str(edges["kegg_edge_subtype"].value_counts().to_dict()),
    )
    check(
        "duplicate classified edges",
        not edges.duplicated(["source", "target", "kegg_edge_subtype"]).any(),
        "no duplicate source-target-subtype rows expected",
    )
    review_count = int(edges["needs_manual_review"].map(bool_value).sum())
    check("manual-review flags", review_count == 0, f"found {review_count}")
    check(
        "primary evidence type",
        set(edges["evidence_type"]) == {"Primary literature"},
        f"types={sorted(set(edges['evidence_type']))}",
    )
    check(
        "source URLs populated",
        edges["source_urls"].str.startswith("http").all(),
        "every retained edge must cite a source",
    )

    actual_set = {edge_tuple(row) for _, row in edges.iterrows()}
    audit_set = {edge_tuple(row) for _, row in audit.iterrows()}
    check(
        "audit register exact match",
        actual_set == audit_set and len(audit) == 32,
        f"actual={len(actual_set)} audit={len(audit_set)}",
    )
    sources_match = True
    for _, row in edges.iterrows():
        key = edge_tuple(row)
        audited_rows = audit[[edge_tuple(r) == key for _, r in audit.iterrows()]]
        if len(audited_rows) != 1:
            sources_match = False
            continue
        primary = clean(audited_rows.iloc[0]["primary_source"])
        if primary not in clean(row["source_urls"]):
            sources_match = False
            errors.append(f"primary source mismatch: {row['source']}→{row['target']}")
    check("primary citations match", sources_match, "each audit citation occurs on its edge")

    forbidden_present = {
        (clean(r["source"]), clean(r["target"]), clean(r["kegg_edge_subtype"]))
        for _, r in edges.iterrows()
    } & FORBIDDEN_EDGES
    check("removed overstatements absent", not forbidden_present, str(sorted(forbidden_present)))

    check(
        "evidence export complete",
        len(evidence) == 32 and evidence["interaction"].str.len().gt(0).all(),
        f"rows={len(evidence)} blank_interactions={(evidence['interaction'].str.len() == 0).sum()}",
    )
    check(
        "evidence review flags",
        "needs_manual_review" in evidence.columns
        and int(evidence["needs_manual_review"].map(bool_value).sum()) == 0,
        "evidence export must preserve resolved review status",
    )

    kgml = OUTPUT / "myrf_pathway.kgml"
    custom_kgml = ROOT / "custom_kegg" / "hsa99999.xml"
    doc = ET.parse(kgml)
    root = doc.getroot()
    entries = root.findall("entry")
    relations = root.findall("relation")
    entry_ids = {entry.attrib["id"] for entry in entries}
    relation_refs = {
        rel.attrib["entry1"] for rel in relations
    } | {
        rel.attrib["entry2"] for rel in relations
    }
    kgml_subtypes = Counter(
        rel.find("subtype").attrib["name"] for rel in relations
    )
    check("KGML entry count", len(entries) == 31, f"found {len(entries)}")
    check("KGML relation count", len(relations) == 32, f"found {len(relations)}")
    check("KGML references resolve", relation_refs <= entry_ids, "all relation IDs must resolve")
    check(
        "KGML subtype parity",
        dict(kgml_subtypes) == EXPECTED_SUBTYPES,
        str(dict(kgml_subtypes)),
    )
    check(
        "Pathview KGML synchronized",
        custom_kgml.exists() and file_hash(kgml) == file_hash(custom_kgml),
        "output and custom_kegg copies must be byte-identical",
    )

    svg = OUTPUT / "myrf_final_pathway.svg"
    svg_root = ET.parse(svg).getroot()
    ns = {"svg": "http://www.w3.org/2000/svg"}
    edge_groups = [
        group for group in svg_root.findall(".//svg:g", ns)
        if group.attrib.get("class") == "edge"
    ]
    check("SVG edge count", len(edge_groups) == 32, f"found {len(edge_groups)}")

    dot = (OUTPUT / "myrf_final_pathway.dot").read_text(encoding="utf-8")
    dot_lines = {
        line.split(" [", 1)[0].strip(): line
        for line in dot.splitlines()
        if " -> " in line and " [" in line
    }
    gpr37_line = dot_lines.get("GPR37 -> MYRF_gene", "")
    hes5_line = dot_lines.get("HES5 -> MYRF_gene", "")
    sox10_line = dot_lines.get("SOX10 -> MYRF_protein", "")
    fbxw7_line = dot_lines.get("FBXW7 -> MYRF_protein", "")
    check(
        "negative indirect rendering",
        all(token in gpr37_line for token in ("arrowhead=tee", "#C0392B", 'style="dashed"'))
        and all(token in hes5_line for token in ("arrowhead=tee", "#C0392B", 'style="dashed"')),
        "GPR37 and HES5 must render as red dashed inhibitory edges",
    )
    check(
        "SOX10 association rendering",
        'dir="none"' in sox10_line and "arrowhead=none" in sox10_line,
        "MYRF-SOX10 association must be visually undirected",
    )
    check(
        "FBXW7 degradation rendering",
        "arrowhead=tee" in fbxw7_line and 'label="+u"' in fbxw7_line,
        "FBXW7 ubiquitination must retain negative/degradative direction",
    )
    wb = openpyxl.load_workbook(WORKBOOK, data_only=False, read_only=True)
    required_sheets = {"Master_Gene_List", "References", "Revision_Log", "Literature_Audit"}
    check(
        "workbook audit sheets",
        required_sheets <= set(wb.sheetnames),
        f"sheets={wb.sheetnames}",
    )
    master = pd.read_excel(WORKBOOK, sheet_name="Master_Gene_List", dtype=str).fillna("")
    mapk1 = master[master.iloc[:, 0] == "MAPK1"]
    check(
        "MAPK1 identifier",
        len(mapk1) == 1 and clean(mapk1.iloc[0, 3]) == "ENSG00000100030",
        "MAPK1 must not reuse the MAPK3 Ensembl ID",
    )
    excluded = master[master.iloc[:, 5] == "Excluded"]
    check("excluded source rows", len(excluded) == 4, f"found {len(excluded)}")
    refs = pd.read_excel(WORKBOOK, sheet_name="References", dtype=str).fillna("")
    check(
        "portable reference URLs",
        not refs.astype(str).apply(lambda c: c.str.contains("libproxy", case=False)).any().any(),
        "References sheet must not contain institution-only proxy URLs",
    )
    workbook_audit = pd.read_excel(WORKBOOK, sheet_name="Literature_Audit", dtype=str).fillna("")
    check(
        "workbook audit row count",
        len(workbook_audit) == 32,
        f"found {len(workbook_audit)}",
    )
    wb.close()

    write_report(checks, edges, audit, errors)
    if errors:
        print(f"Round-2 audit FAILED with {len(errors)} issue(s).")
        for error in errors:
            print(f"  - {error}")
        print(f"Report: {REPORT}")
        return 1

    print("Round-2 audit PASSED.")
    print(f"Validated {len(nodes)} nodes and {len(edges)} literature-reviewed edges.")
    print(f"Report: {REPORT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
